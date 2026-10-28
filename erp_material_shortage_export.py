"""Excel da fotografia exibida em Faltas por chassi; não recalcula nem grava dados."""
from collections import OrderedDict
from datetime import date, datetime
from io import BytesIO
import math
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE


COLUMNS = [
    ("Seq.", "sequencia", 8), ("Entrega do veículo", "data_entrega", 18),
    ("O.S.", "numero_os", 12), ("Cliente", "cliente_nome", 28),
    ("SKU da composição", "codigo", 20), ("Descrição", "descricao", 65),
    ("Un.", "unidade", 8), ("Necessário", "necessario", 16),
    ("Atendido na O.S.", "atendido_na_os_total", 18), ("Saldo físico", "saldo_atual", 16),
    ("Empenhado global", "saldo_empenhado", 19), ("Reserva anterior", "reserva", 18),
    ("Disponível na vez", "saldo_disponivel_sequencia", 19), ("Falta", "faltante", 16),
    ("Em trânsito*", "em_chegada", 16), ("Previsão de chegada", "previsao_chegada", 20),
]
NOTE = ("Saldo disponível na vez = saldo físico − empenhado global − reserva dos cards anteriores. "
        "A falta já desconta o atendimento próprio desta O.S.\n"
        "* Em trânsito é previsão informativa; não foi abatida da falta atual.")
BANK_NOTE = ("Sem conversão automática de bancos: o relatório mantém o SKU exato da composição da O.S. "
             "PME, BJD e C/TRILHO não são substituídos.")
SEARCH_FIELDS = ("chassi", "numero_os", "codigo", "descricao", "cliente_nome", "proposta_numero", "veiculo", "transformacao")


def _text(cell, value):
    cell.value = ILLEGAL_CHARACTERS_RE.sub("", str(value if value is not None else "—"))
    cell.data_type = "s"  # Descrições e códigos nunca viram fórmulas.


def _number(value):
    try:
        result = float(value)
        return result if math.isfinite(result) else None
    except (ValueError, TypeError):
        return None


def _sequence(order):
    value = _number(order.get("sequencia"))
    return value if value is not None else math.inf


def _matches(row, query):
    return not query or query in " ".join(str(row.get(key) or "") for key in SEARCH_FIELDS).lower()


def _date_text(value, timestamp=False):
    try:
        return datetime.fromisoformat(str(value)).strftime("%d/%m/%Y %H:%M:%S" if timestamp else "%d/%m/%Y")
    except (ValueError, TypeError):
        return "—"


def _band(ws, text, fill="EDF4FA", color="0F2744", height=30):
    row = ws.max_row + 1 if ws.cell(1, 1).value is not None else 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=16)
    cell = ws.cell(row, 1)
    _text(cell, text)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(name="Arial", size=11, bold=True, color=color)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = height
    return row


def _table(ws, rows):
    if not rows:
        _band(ws, "Nenhum item em falta para este filtro.", "F8FAFC", "64748B")
        return
    header = ws.max_row + 1
    if not getattr(ws, "_shortage_header_set", False):
        ws.freeze_panes = f"F{header + 1}"
        ws.print_title_rows = f"1:{header}"
        ws._shortage_header_set = True
    for col, (label, _, _) in enumerate(COLUMNS, 1):
        cell = ws.cell(header, col, label)
        cell.fill = PatternFill("solid", fgColor="EFF4F8")
        cell.font = Font(name="Arial", size=10, bold=True, color="475569")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[header].height = 30
    for row in rows:
        r = ws.max_row + 1
        for col, (_, key, _) in enumerate(COLUMNS, 1):
            cell = ws.cell(r, col)
            value = row.get(key)
            cell.font = Font(name="Arial", size=11, color="172033", bold=col in (3, 5))
            cell.alignment = Alignment(vertical="top", wrap_text=col in (4, 6))
            if col in (2, 16):
                try:
                    cell.value = date.fromisoformat(str(value)[:10])
                    cell.number_format = "dd/mm/yyyy"
                except (ValueError, TypeError):
                    _text(cell, "—")
            elif col == 1 or 8 <= col <= 15:
                cell.value = _number(value) if col == 1 else (_number(value) or 0)
                cell.number_format = "#,##0.###"
                cell.alignment = Alignment(horizontal="right", vertical="top")
            else:
                if key == "descricao" and any(s != row.get("codigo") for s in (row.get("sku_origens") or [])):
                    value = str(value or "") + "\nSKU de origem: " + ", ".join(row["sku_origens"])
                _text(cell, value or "—")
            if col == 14:
                cell.fill = PatternFill("solid", fgColor="FFF6F5")
                cell.font = Font(name="Arial", size=11, color="B42318", bold=True)
        ws.row_dimensions[r].height = max(30, 15 * (1 + len(str(row.get("descricao") or "")) // 60))
    if ws.title == "Consolidado":
        ws.auto_filter.ref = f"A{header}:P{ws.max_row}"
    _band(ws, NOTE, "F8FAFC", "475569", 35)


def build_workbook(report, query="", active_key="__ALL__"):
    if not isinstance(report, dict) or not isinstance(report.get("rows"), list) or not isinstance(report.get("chassis"), list):
        raise ValueError("Fotografia do relatório inválida. Recalcule a tela e tente novamente.")
    orders = report["chassis"]
    if any(not isinstance(o, dict) or not isinstance(o.get("items", []), list) for o in orders):
        raise ValueError("Composição da fotografia inválida.")
    if len(report["rows"]) + sum(len(o.get("items", [])) for o in orders) > 200000:
        raise ValueError("Relatório muito grande. Reduza o escopo antes de exportar.")
    query = str(query or "").strip().lower()
    wb = Workbook()
    wb.remove(wb.active)
    groups = OrderedDict()
    for order in orders:
        groups.setdefault(str(order.get("chassi") or "Sem chassi"), []).append(order)
    groups = sorted(groups.items(), key=lambda g: (
        min(_sequence(o) for o in g[1]),
        min((_number(o.get("item_number")) or 0) for o in g[1]), g[0],
    ))

    def sheet(name, title):
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        for c, (_, _, width) in enumerate(COLUMNS, 1):
            ws.column_dimensions[get_column_letter(c)].width = width
        _band(ws, title, "0F2744", "FFFFFF", 34)
        _band(ws, f"O.S. no WIP: {report.get('wip_count', 0)}    Chassis: {report.get('chassi_count', 0)}    "
              f"Linhas com falta: {report.get('shortage_lines', 0)}    SKUs com falta: {report.get('shortage_skus', 0)}")
        _band(ws, f"Atualizado em: {_date_text(report.get('gerado_em'), True)}    Filtro: {query or 'Sem filtro'}", "F8FAFC", "64748B")
        if report.get("incomplete_count"):
            _band(ws, f"Atenção: {report['incomplete_count']} O.S. da sequência sem composição válida. Confira os registros sinalizados.", "FFF7DF", "854D0E")
        ws.freeze_panes = f"F{ws.max_row + 2}"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        return ws

    ws = sheet("Consolidado", "Faltas de materiais por chassi — Consolidado")
    _table(ws, [r for r in report["rows"] if _matches(r, query)])
    _band(ws, BANK_NOTE, "F8FAFC", "475569", 32)
    for chassi, chassis_orders in groups:
        items = sorted({str(o.get("item_number")) for o in chassis_orders if o.get("item_number") is not None}, key=lambda v: (_number(v) or 0, v))
        suffix = re.sub(r"[^A-Za-z0-9]", "", chassi)[-8:] or "SEM CHASSI"
        item_label = re.sub(r"[\\/*?:\[\]]", "-", '-'.join(items) or 'SEM ITEM')
        base = f"{item_label[:30-len(suffix)]} {suffix}"
        name, counter = base, 1
        while name.lower() in {n.lower() for n in wb.sheetnames}:
            counter += 1
            name = f"{base[:26]} ({counter})"[:31]
        ws = sheet(name, f"Faltas de materiais por chassi — {chassi}")
        if chassi == active_key:
            wb.active = wb.sheetnames.index(name)
        for order in chassis_orders:
            _band(ws, f"O.S. {order.get('numero_os', '—')} · {chassi}")
            _band(ws, f"Item: {order.get('item_number', '—')}    Sequência: {order.get('sequencia', '—')}    "
                  f"Entrega: {_date_text(order.get('data_entrega'))}    Situação: {str(order.get('status') or '').replace('_', ' ')}", "EDF4FA", "475569")
            _band(ws, f"Cliente: {order.get('cliente_nome') or '—'}    Proposta: {order.get('proposta_numero') or '—'}\n"
                  f"Veículo: {' / '.join(str(order.get(k)) for k in ('veiculo', 'transformacao') if order.get(k)) or '—'}", height=36)
            if not order.get("documento_encontrado"):
                _band(ws, "Composição da O.S. não encontrada: não é possível calcular suas necessidades.", "FFF7DF", "854D0E")
            elif not (order.get("reserva_sequencia") or {}).get("completa"):
                _band(ws, "A sequência possui composição anterior incompleta; saldo/reserva exigem conferência.", "FFF7DF", "854D0E")
            rows = [{**i, **order} for i in order.get("items", []) if (_number(i.get("faltante")) or 0) > 0 and _matches({**i, **order}, query)]
            _table(ws, rows)
        _band(ws, BANK_NOTE, "F8FAFC", "475569", 32)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
