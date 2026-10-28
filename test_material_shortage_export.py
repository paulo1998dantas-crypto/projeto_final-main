import copy
from datetime import datetime
import unittest

from openpyxl import load_workbook

from erp_material_shortage_export import COLUMNS, build_workbook


def fixture():
    orders = []
    for seq, item, vin in [(12, 3186, 'WF0LTBVF8TU022142'), (3, 3174, '9VYF62001SJ064758')]:
        line = dict(codigo='00123456', descricao='Banco normal', unidade='pc', necessario=6.5,
                    atendido_na_os_total=1, saldo_atual=8, saldo_empenhado=3, reserva=4,
                    saldo_disponivel_sequencia=1, faltante=4.5, em_chegada=2,
                    previsao_chegada='2026-11-03')
        order = dict(sequencia=seq, item_number=item, chassi=vin, numero_os=str(item),
                     cliente_nome='Cliente', proposta_numero='10060', data_entrega='2026-11-04',
                     status='EM_PRODUÇÃO', documento_encontrado=True, reserva_sequencia={'completa': True})
        order['items'] = [{**line, **order}]
        orders.append(order)
    return dict(gerado_em='2026-10-28T10:00:00-03:00', wip_count=2, chassi_count=2,
                shortage_lines=2, shortage_skus=1, incomplete_count=0, chassis=orders,
                rows=[orders[1]['items'][0], orders[0]['items'][0]])


def data_rows(ws):
    return [row for row in ws.iter_rows() if row[4].value == '00123456']


class ExportTests(unittest.TestCase):
    def test_snapshot_layout_typed_values_and_sequence(self):
        report = fixture()
        before = copy.deepcopy(report)
        wb = load_workbook(build_workbook(report, active_key=report['chassis'][0]['chassi']))
        self.assertEqual(report, before)
        self.assertEqual(wb.sheetnames, ['Consolidado', '3174 SJ064758', '3186 TU022142'])
        self.assertEqual(wb.active.title, '3186 TU022142')
        sheet = wb['Consolidado']
        self.assertEqual([c.value for c in sheet[4]], [c[0] for c in COLUMNS])
        rows = data_rows(sheet)
        self.assertEqual(len(rows), 2)
        self.assertIsInstance(rows[0][1].value, datetime)
        self.assertEqual(rows[0][4].data_type, 's')
        self.assertEqual(rows[0][7].value, 6.5)
        self.assertEqual(rows[0][11].value, 4)
        self.assertEqual(rows[0][13].value, 4.5)
        self.assertEqual(rows[0][13].fill.fgColor.rgb[-6:], 'FFF6F5')
        self.assertEqual(sheet.freeze_panes, 'F5')
        self.assertEqual(sheet.auto_filter.ref, 'A4:P6')
        self.assertTrue(data_rows(wb.worksheets[1]))
        self.assertEqual(wb.worksheets[1].freeze_panes, 'F8')

    def test_filter_preserves_tabs_and_same_lines_as_digital(self):
        report = fixture()
        wb = load_workbook(build_workbook(report, query='SJ064758'))
        self.assertEqual(len(data_rows(wb.worksheets[0])), 1)
        self.assertEqual(len(data_rows(wb.worksheets[1])), 1)
        self.assertEqual(len(data_rows(wb.worksheets[2])), 0)

    def test_formula_like_text_does_not_become_formula(self):
        report = fixture()
        report['rows'][0]['descricao'] = '=HYPERLINK("x")'
        wb = load_workbook(build_workbook(report))
        self.assertEqual(wb.worksheets[0]['F5'].data_type, 's')
        self.assertEqual(wb.worksheets[0]['F5'].value, '=HYPERLINK("x")')

    def test_repeated_chassis_stays_one_sheet_with_separate_orders(self):
        report = fixture()
        report['chassis'][0]['chassi'] = report['chassis'][1]['chassi']
        wb = load_workbook(build_workbook(report))
        self.assertEqual(len(wb.worksheets), 2)
        self.assertEqual(wb.worksheets[1].title, '3174-3186 SJ064758')
        self.assertEqual(len(data_rows(wb.worksheets[1])), 2)

    def test_empty_and_incomplete_reports_remain_explicit(self):
        wb = load_workbook(build_workbook({'rows': [], 'chassis': [], 'incomplete_count': 1}))
        self.assertEqual(wb.sheetnames, ['Consolidado'])
        values = [c.value for row in wb.active for c in row if c.value]
        self.assertTrue(any('sem composição válida' in str(v) for v in values))
        self.assertTrue(any('Nenhum item em falta' in str(v) for v in values))

    def test_invalid_snapshot_is_rejected(self):
        for value in [None, {}, {'rows': [], 'chassis': [None]}]:
            with self.assertRaises(ValueError):
                build_workbook(value)


if __name__ == '__main__':
    unittest.main()
