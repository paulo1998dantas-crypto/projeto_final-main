import unittest
from unittest.mock import patch
from datetime import date
from uuid import uuid4
from types import SimpleNamespace
from io import BytesIO
from sqlalchemy import create_engine, text, select, func
from sqlalchemy.pool import StaticPool
from fastapi import FastAPI
from fastapi.testclient import TestClient
from fastapi.templating import Jinja2Templates
from openpyxl import load_workbook
import quality_ftt as q
import quality_routes
import authz

class FTTTests(unittest.TestCase):
    def setUp(self):
        self.engine=create_engine("sqlite://",connect_args={"check_same_thread":False},poolclass=StaticPool)
        q.metadata.create_all(self.engine)
        self.work=str(uuid4())
        self.user=authz.Principal(9,"PAULO","PAULO",True,1,frozenset({"QUALIDADE"}),
                                authz._default_permissions({"QUALIDADE"}))
        with self.engine.begin() as c:
            for sql in [
                "create table erp_vehicles(id text primary key,chassi text)",
                "create table erp_vehicle_entries(id text primary key,vehicle_id text,item_number integer)",
                "create table erp_work_orders(id text primary key,numero_os text,vehicle_entry_id text,status text,is_current boolean)",
                "create table erp_work_order_stages(work_order_id text,stage_code text,status text,aplicavel boolean,termino timestamp)"
            ]: c.execute(text(sql))
            c.execute(text("insert into erp_vehicles values ('v','CHASSI-REPETIDO')"))
            c.execute(text("insert into erp_vehicle_entries values ('e','v',3185)"))
            c.execute(text("insert into erp_work_orders values (:id,'3185','e','ATIVA',true)"),{"id":self.work})
            c.execute(text("insert into erp_work_order_stages values (:id,'BCO','CONCLUÍDA',true,'2026-10-27T15:00:00+00:00')"),{"id":self.work})
            c.execute(text("insert into erp_work_order_stages values (:id,'ELÉTRICA','PENDENTE',true,null)"),{"id":self.work})
            self.defect=q.create_defect(c,"Tecido enrugado","PAULO")["id"]
        self.clock=patch("quality_ftt.today",return_value=date(2026,11,1));self.clock.start()
    def tearDown(self):
        self.clock.stop();self.engine.dispose()
    def payload(self,**kw):
        return dict(idempotency_key=str(uuid4()),work_order_id=self.work,observed_on="2026-10-27",
            issues=[dict(defect_id=self.defect,stage_code="BCO",disposition="DEFEITO")],**kw)
    def save(self,p):
        with self.engine.begin() as c:return q.create_report(c,p,self.user,["BCO","ELÉTRICA"])
    def report(self):
        with self.engine.connect() as c:return q.dashboard(c,["BCO","ELÉTRICA"],"2026-10-26","2026-10-31")
    def test_two_defects_same_os_count_one_bad_unit_and_two_occurrences(self):
        p=self.payload()
        with self.engine.begin() as c:
            another=q.create_defect(c,"Banco desalinhado","PAULO")["id"]
        p["issues"].append(dict(defect_id=another,stage_code="BCO",disposition="DEFEITO"))
        self.save(p);r=self.report()
        self.assertEqual(2,r["total_failures"])
        self.assertEqual(1,r["areas"][0]["produced"])
        self.assertEqual(1,r["areas"][0]["defective"])
        self.assertEqual(0,r["areas"][0]["ftt"])
        self.assertIsNone(r["areas"][1]["ftt"])
    def test_idempotent_retry_and_different_payload_rejected(self):
        p=self.payload();self.save(p);self.assertTrue(self.save(p)["replayed"])
        p["issues"][0]["disposition"]="REFUGO"
        self.assertTrue(self.save(p)["replayed"])
        p["issues"][0]["stage_code"]="ELÉTRICA"
        with self.assertRaises(ValueError):self.save(p)
        self.assertEqual(1,self.report()["total_failures"])
    def test_multiple_defect_reports_reject_the_same_unit_once(self):
        self.save(self.payload())
        p=self.payload();p["issues"][0]["disposition"]="REFUGO";self.save(p)
        area=self.report()["areas"][0]
        self.assertEqual((1,0),(area["defective"],area["first_pass"]))

    def test_new_report_forces_defect_even_when_a_legacy_value_is_sent(self):
        p=self.payload();p["issues"][0]["disposition"]="REFUGO";self.save(p)
        with self.engine.connect() as c:
            self.assertEqual("DEFEITO",c.execute(select(q.issues.c.disposition)).scalar_one())
    def test_pending_area_occurrences_only_enter_ftt_after_completion(self):
        p=self.payload();p["issues"][0]["stage_code"]="ELÉTRICA";self.save(p)
        self.assertEqual(1,self.report()["pending"])
        with self.engine.begin() as c:
            c.execute(text("update erp_work_order_stages set status='CONCLUÍDA',termino='2026-10-29T15:00:00+00:00' where stage_code='ELÉTRICA'"))
        r=self.report();self.assertEqual(0,r["pending"]);self.assertEqual(1,r["areas"][1]["defective"])
    def test_edited_completion_date_is_used(self):
        self.save(self.payload())
        with self.engine.begin() as c:
            c.execute(text("update erp_work_order_stages set termino='2026-11-01T15:00:00+00:00' where stage_code='BCO'"))
        self.assertEqual(0,self.report()["areas"][0]["produced"])
    def test_cutoff_and_local_timezone(self):
        completions=[dict(work_order_id="old",stage_code="BCO",termino="2026-10-26T02:00:00+00:00"),
                     dict(work_order_id="new",stage_code="BCO",termino="2026-10-26T03:00:00+00:00")]
        r=q.calculate(completions,[],["BCO"],q.START,date(2026,10,31))
        self.assertEqual(1,r["totals"]["produced"])
        self.assertEqual(100,r["totals"]["ftt"])
        self.assertEqual(q.START,q.day("2026-10-26"))
    def test_pareto_top_ten_uses_all_occurrences_and_groups_same_defect_across_areas(self):
        rows=[dict(description=str(i),stage_code="BCO",work_order_id="w",observed_on="2026-10-27",disposition="DEFEITO") for i in range(11)]
        rows.append(dict(rows[0],stage_code="ELÉTRICA"))
        r=q.calculate([],rows,["BCO","ELÉTRICA"],q.START,date(2026,10,31))
        self.assertEqual(10,len(r["pareto"]));self.assertEqual(12,r["total_failures"])
        self.assertEqual(2,r["pareto"][0]["frequency"])
        self.assertLess(r["pareto"][-1]["accumulated"],100)
    def test_work_order_revision_keeps_the_same_vehicle_quality_failure(self):
        completions=[dict(work_order_id="revision-2",unit_id="entry",stage_code="BCO",termino="2026-10-27T15:00:00+00:00")]
        occurrences=[dict(work_order_id="revision-1",unit_id="entry",stage_code="BCO",description="Falha",
                          observed_on="2026-10-27",disposition="DEFEITO")]
        r=q.calculate(completions,occurrences,["BCO"],q.START,date(2026,10,31))
        self.assertEqual(1,r["areas"][0]["defective"])
        self.assertEqual(0,r["areas"][0]["ftt"])
    def test_invalid_line_rolls_back_entire_report(self):
        p=self.payload();p["issues"].append(dict(defect_id=self.defect,stage_code="INVALID",disposition="DEFEITO"))
        with self.assertRaises(ValueError):self.save(p)
        self.assertEqual(0,self.report()["total_failures"])
    def test_dates_duplicate_lines_and_not_applicable_rejected(self):
        p=self.payload()
        for when in ("2026-10-25","2026-11-02"):
            p["observed_on"]=when
            with self.assertRaises(ValueError):self.save(p)
        p["observed_on"]="2026-10-27";p["issues"].append(p["issues"][0])
        with self.assertRaises(ValueError):self.save(p)
        with self.engine.begin() as c:c.execute(text("update erp_work_order_stages set aplicavel=false"))
        p["issues"]=p["issues"][:1]
        with self.assertRaises(ValueError):self.save(p)
    def test_normalized_defect_and_same_chassis_distinct_work_orders(self):
        with self.engine.begin() as c:
            self.assertEqual(self.defect,q.create_defect(c,"  TECIDO   ENRUGADO  ","TEST")["id"])
        rows=[dict(work_order_id=w,stage_code="BCO",termino="2026-10-27T15:00:00+00:00") for w in ["one","two"]]
        self.assertEqual(2,q.calculate(rows,[],["BCO"],q.START,date(2026,10,31))["totals"]["produced"])
    def client(self):
        app=FastAPI()
        class DB:
            def close(self):pass
        quality_routes.register(app,self.engine,DB,lambda request,db:self.user,
            Jinja2Templates(directory="templates"),[("BCO",1,[]),("ELÉTRICA",2,[])])
        return TestClient(app)
    def test_http_permissions_csrf_template_and_excel(self):
        client=self.client()
        self.assertEqual(200,client.get("/ftt").status_code)
        self.assertEqual(403,client.post("/api/ftt/reports",json=self.payload()).status_code)
        response=client.post("/api/ftt/reports",json=self.payload(),headers={"X-FTT-Request":"1","Origin":"http://testserver"})
        self.assertEqual(200,response.status_code,response.text)
        book=load_workbook(BytesIO(client.get("/api/ftt/export.xlsx?from=2026-10-26&to=2026-10-31").content))
        self.assertEqual(["Pareto","FTT por área","Apontamentos"],book.sheetnames)
        self.user=authz.Principal(8,"OPERADOR","OPERADOR",True,1,frozenset({"APONTAMENTO"}),
                                 authz._default_permissions({"APONTAMENTO"}))
        self.assertEqual(403,client.get("/api/ftt").status_code)
        self.assertEqual(403,client.post("/api/ftt/reports",json=self.payload(),headers={"X-FTT-Request":"1"}).status_code)
    def test_quality_role_does_not_gain_production_writes(self):
        self.assertTrue(self.user.can("mes.quality.write"))
        self.assertFalse(self.user.can("mes.stage.write"))
        self.assertFalse(self.user.can("mes.work_orders.manage"))

    def test_release_is_never_offered_as_an_ftt_area(self):
        app=FastAPI()
        class DB:
            def close(self):pass
        quality_routes.register(app,self.engine,DB,lambda request,db:self.user,
            Jinja2Templates(directory="templates"),[("BCO",1,[]),("LIBERAÇÃO",2,[])])
        options=TestClient(app).get("/api/ftt/options").json()
        self.assertEqual(["BCO"],options["stages"])

if __name__=="__main__":unittest.main()

